from sqlalchemy.orm import Session
from fastapi import HTTPException, status
from typing import List, Optional
from datetime import datetime, date
import os
import uuid
from docxtpl import DocxTemplate
from openpyxl import Workbook

from app.models import ReportTemplate, GeneratedReport, Student, Group, User, Competition
from app.schemas import ReportTemplateCreate


# ========== ШАБЛОНЫ ==========

def get_templates(
        db: Session,
        college_id: int,
        report_type: Optional[str] = None
) -> List[ReportTemplate]:
    query = db.query(ReportTemplate).filter(
        ReportTemplate.college_id == college_id,
        ReportTemplate.is_active == True
    )
    if report_type:
        query = query.filter(ReportTemplate.report_type == report_type)
    return query.all()


def get_template_by_id(db: Session, template_id: int) -> ReportTemplate:
    template = db.query(ReportTemplate).filter(ReportTemplate.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    return template


def create_template(db: Session, template_in, college_id: int) -> ReportTemplate:
    template = ReportTemplate(
        college_id=college_id,
        name=template_in.name,
        description=template_in.description,
        report_type=template_in.report_type,
        template_file_path=template_in.template_file_path,
        template_file_name=template_in.template_file_name,
        template_type=template_in.template_type,
        placeholders=template_in.placeholders,
        is_active=True
    )
    db.add(template)
    db.commit()
    db.refresh(template)
    return template


def delete_template(db: Session, template_id: int, college_id: int) -> None:
    template = db.query(ReportTemplate).filter(
        ReportTemplate.id == template_id,
        ReportTemplate.college_id == college_id
    ).first()
    if not template:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    template.is_active = False
    db.commit()


# ========== ГЕНЕРАЦИЯ ОТЧЁТОВ ==========

def generate_report(
        db: Session,
        template_id: int,
        group_id: Optional[int],
        start_date: Optional[date],
        end_date: Optional[date],
        format: str,
        curator_id: int,
        college_id: int
) -> dict:
    template = get_template_by_id(db, template_id)

    # Собираем данные для отчёта
    context = {}

    if template.report_type == "student_list" and group_id:
        students = db.query(Student).filter(
            Student.college_id == college_id,
            Student.is_active == True
        ).join(Student.group_students).filter(Student.group_students.any(group_id=group_id))

        context = {
            "students": [
                {
                    "full_name": s.user.full_name if s.user else "",
                    "birth_date": s.birth_date.strftime("%d.%m.%Y") if s.birth_date else "",
                    "email": s.email or "",
                    "phone": s.phone or "",
                    "personal_number": s.personal_number
                }
                for s in students
            ]
        }

    # Генерация файла
    reports_dir = "uploads/reports"
    os.makedirs(reports_dir, exist_ok=True)

    filename = f"report_{uuid.uuid4().hex[:8]}.{format}"
    file_path = os.path.join(reports_dir, filename)

    if format == "pdf":
        # Генерация PDF через WeasyPrint
        from weasyprint import HTML
        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; }}
                table {{ width: 100%; border-collapse: collapse; }}
                th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                th {{ background-color: #f2f2f2; }}
            </style>
        </head>
        <body>
            <h1>Отчёт: {template.name}</h1>
            <p>Дата: {datetime.now().strftime("%d.%m.%Y %H:%M")}</p>
            <table>
                <thead>
                    <tr>
                        <th>№</th>
                        <th>ФИО</th>
                        <th>Дата рождения</th>
                        <th>Email</th>
                        <th>Телефон</th>
                    </tr>
                </thead>
                <tbody>
                    {''.join([
            f"<tr><td>{i + 1}</td><td>{s['full_name']}</td><td>{s['birth_date']}</td><td>{s['email']}</td><td>{s['phone']}</td></tr>"
            for i, s in enumerate(context.get('students', []))
        ])}
                </tbody>
            </table>
        </body>
        </html>
        """
        HTML(string=html_content).write_pdf(file_path)
        media_type = "application/pdf"

    elif format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        # Заголовки
        headers = ["№", "ФИО", "Дата рождения", "Email", "Телефон"]
        ws.append(headers)

        # Данные
        for i, s in enumerate(context.get('students', [])):
            ws.append([
                i + 1,
                s.get('full_name', ''),
                s.get('birth_date', ''),
                s.get('email', ''),
                s.get('phone', '')
            ])

        wb.save(file_path)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise HTTPException(status_code=400, detail=f"Формат {format} не поддерживается")

    # Сохраняем в базу
    generated = GeneratedReport(
        curator_id=curator_id,
        template_id=template_id,
        academic_year_id=None,  # можно добавить логику
        file_path=file_path,
        file_format=format,
        params={"group_id": group_id, "start_date": str(start_date), "end_date": str(end_date)}
    )
    db.add(generated)
    db.commit()
    db.refresh(generated)

    return {
        "file_path": file_path,
        "media_type": media_type,
        "filename": filename
    }


def get_generated_reports(db: Session, curator_id: int, skip: int = 0, limit: int = 15) -> List[dict]:
    reports = db.query(GeneratedReport).filter(
        GeneratedReport.curator_id == curator_id
    ).offset(skip).limit(limit).all()

    return [
        {
            "id": r.id,
            "template_name": r.template.name if r.template else "",
            "file_format": r.file_format,
            "generated_at": r.generated_at,
            "file_path": r.file_path
        }
        for r in reports
    ]


def get_generated_report_by_id(db: Session, report_id: int, curator_id: int) -> dict:
    report = db.query(GeneratedReport).filter(
        GeneratedReport.id == report_id,
        GeneratedReport.curator_id == curator_id
    ).first()
    if not report:
        raise HTTPException(status_code=404, detail="Отчёт не найден")

    return {
        "file_path": report.file_path,
        "media_type": "application/pdf" if report.file_format == "pdf" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "filename": f"report_{report.id}.{report.file_format}"
    }


def delete_generated_report(db: Session, report_id: int, curator_id: int) -> None:
    report = db.query(GeneratedReport).filter(
        GeneratedReport.id == report_id,
        GeneratedReport.curator_id == curator_id
    ).first()
    if not report:
        raise HTTPException(status_code=404, detail="Отчёт не найден")

    # Удаляем файл с диска
    if os.path.exists(report.file_path):
        os.remove(report.file_path)

    db.delete(report)
    db.commit()